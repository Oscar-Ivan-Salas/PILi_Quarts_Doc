from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def set_style(paragraph, font_size=11, bold=False, color=(0, 0, 0), align=None):
    if align:
        paragraph.alignment = align
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)

def add_tesla_header(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "TESLA ELECTRICIDAD Y AUTOMATIZACIÓN S.A.C."
    set_style(p, font_size=14, bold=True, color=(0, 82, 163), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    p2 = header.add_paragraph("RUC: 20601138787 | Jr. Las Ágatas Mz B Lote 09, Urb. San Carlos, SJL")
    set_style(p2, font_size=9, color=(75, 85, 99), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    p3 = header.add_paragraph("Email: ingenieria.teslaelectricidad@gmail.com | Tel: 906 315 961")
    set_style(p3, font_size=9, color=(75, 85, 99), align=WD_ALIGN_PARAGRAPH.CENTER)

def add_tesla_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Documento Generado por Sistema RALFTH - Tesla Electricidad v11.0"
    set_style(p, font_size=8, color=(156, 163, 175), align=WD_ALIGN_PARAGRAPH.CENTER)

def create_cotizacion_master(path, type_name="COMPLEJA"):
    doc = Document()
    add_tesla_header(doc)
    add_tesla_footer(doc)
    
    # Title
    doc.add_paragraph("\n")
    p = doc.add_paragraph(f"COTIZACIÓN PROFESIONAL {type_name}")
    set_style(p, font_size=20, bold=True, color=(0, 82, 163), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    p_code = doc.add_paragraph("N° {{ numero_cotizacion }}")
    set_style(p_code, font_size=12, bold=True, color=(30, 64, 175), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    doc.add_paragraph("\n")
    
    # Info Table
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    
    cells = [
        ("CLIENTE:", "{{ cliente_nombre }}"),
        ("PROYECTO:", "{{ proyecto_nombre }}"),
        ("FECHA:", "{{ fecha }}"),
        ("VALIDEZ:", "30 días calendario")
    ]
    
    for i, (label, value) in enumerate(cells):
        table.cell(i,0).text = label
        table.cell(i,1).text = value
        table.cell(i,0).paragraphs[0].runs[0].font.bold = True
    
    doc.add_paragraph("\n")
    
    # Items Table
    doc.add_heading("DETALLE TÉCNICO Y ECONÓMICO", level=2)
    table_items = doc.add_table(rows=1, cols=6)
    table_items.style = 'Table Grid'
    hdr = table_items.rows[0].cells
    hdr[0].text = "ITEM"
    hdr[1].text = "DESCRIPCIÓN"
    hdr[2].text = "UND"
    hdr[3].text = "CANT"
    hdr[4].text = "P.UNIT"
    hdr[5].text = "TOTAL"
    
    # Loop row
    row = table_items.add_row().cells
    row[0].text = "{{ item.item }}"
    row[1].text = "{{ item.descripcion }}"
    row[2].text = "{{ item.unidad }}"
    row[3].text = "{{ item.cantidad }}"
    row[4].text = "{{ item.precio_unitario }}"
    row[5].text = "{{ item.total }}"
    
    # The tr loop tags for docxtpl
    # Note: docxtpl prefers tags in the cells or as comments. 
    # For simplicity in this script, we'll put them in the text.
    p0 = row[0].paragraphs[0]
    p0.text = ""
    p0.add_run("{% tr for item in items %}")
    p0.add_run("{{ item.item }}")
    
    row[5].paragraphs[0].add_run("{% tr endfor %}")

    doc.add_paragraph("\n")
    
    # Totals
    p_tot = doc.add_paragraph()
    p_tot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tot.add_run("SUBTOTAL: {{ subtotal }}\n").bold = True
    p_tot.add_run("IGV (18%): {{ igv }}\n").bold = True
    p_tot.add_run("TOTAL: {{ total }}").bold = True
    
    doc.save(path)

def create_informe_apa_master(path):
    doc = Document()
    add_tesla_header(doc)
    add_tesla_footer(doc)
    
    # Portada
    doc.add_paragraph("\n" * 4)
    p = doc.add_paragraph("INFORME EJECUTIVO APA")
    set_style(p, font_size=16, bold=True, color=(0, 0, 0), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    p2 = doc.add_paragraph("{{ titulo_proyecto }}")
    set_style(p2, font_size=14, bold=False, color=(0, 0, 0), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    doc.add_paragraph("\n" * 4)
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.add_run("Cliente: {{ cliente }}\n")
    p_info.add_run("Fecha: {{ fecha }}\n")
    p_info.add_run("Código: {{ codigo_informe }}")
    
    doc.add_page_break()
    
    # Content
    doc.add_heading("1. RESUMEN EJECUTIVO", level=1)
    doc.add_paragraph("{{ resumen_ejecutivo }}")
    
    doc.add_heading("2. KPIs FINANCIEROS", level=1)
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    table.cell(0,0).text = "Inversión:"
    table.cell(0,1).text = "$ {{ presupuesto }}"
    table.cell(1,0).text = "ROI:"
    table.cell(1,1).text = "{{ roi_estimado }}%"
    table.cell(2,0).text = "Payback:"
    table.cell(2,1).text = "{{ payback_meses }} meses"
    table.cell(3,0).text = "TIR:"
    table.cell(3,1).text = "{{ tir_proyectada }}%"
    
    doc.add_heading("3. NORMATIVA", level=1)
    doc.add_paragraph("El proyecto se rige bajo la norma {{ normativa_aplicable }}.")
    
    doc.save(path)

def create_informe_tecnico_master(path):
    doc = Document()
    add_tesla_header(doc)
    add_tesla_footer(doc)
    
    doc.add_heading("INFORME TÉCNICO: {{ titulo_informe }}", level=0)
    
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    table.cell(0,0).text = "CÓDIGO:"
    table.cell(0,1).text = "{{ codigo_informe }}"
    table.cell(1,0).text = "CLIENTE:"
    table.cell(1,1).text = "{{ cliente }}"
    table.cell(2,0).text = "FECHA:"
    table.cell(2,1).text = "{{ fecha }}"
    
    doc.add_heading("RESUMEN", level=1)
    doc.add_paragraph("{{ resumen_ejecutivo }}")
    
    doc.add_heading("ALCANCE TÉCNICO", level=1)
    doc.add_paragraph("Servicio: {{ servicio_nombre }}")
    doc.add_paragraph("Norma: {{ normativa_aplicable }}")
    
    doc.save(path)

def create_proyecto_pmi_master(path, type_name="COMPLEJO"):
    doc = Document()
    add_tesla_header(doc)
    add_tesla_footer(doc)
    
    doc.add_heading(f"PROJECT CHARTER {type_name}", level=0)
    doc.add_paragraph("{{ nombre_proyecto }}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    table.cell(0,0).text = "Código:"
    table.cell(0,1).text = "{{ codigo_proyecto }}"
    table.cell(1,0).text = "Cliente:"
    table.cell(1,1).text = "{{ cliente }}"
    table.cell(2,0).text = "Inicio:"
    table.cell(2,1).text = "{{ fecha_inicio }}"
    table.cell(3,0).text = "Fin:"
    table.cell(3,1).text = "{{ fecha_fin }}"
    
    doc.add_heading("KPIs DE CONTROL", level=1)
    doc.add_paragraph("Presupuesto: $ {{ presupuesto }}")
    doc.add_paragraph("SPI: {{ spi }} | CPI: {{ cpi }}")
    doc.add_paragraph("EV: {{ ev_k }}K | PV: {{ pv_k }}K | AC: {{ ac_k }}K")
    
    doc.add_heading("ALCANCE", level=1)
    doc.add_paragraph("{{ alcance_proyecto }}")
    
    doc.save(path)

# Ensure directory exists
output_dir = r"e:\PILi_Quarts\PILi_Quarts_V3.0\backend\modules\N04_Binary_Factory\templates"
os.makedirs(output_dir, exist_ok=True)

# Generate all 6
create_cotizacion_master(os.path.join(output_dir, "master_cotizacion_compleja.docx"), "COMPLEJA")
create_cotizacion_master(os.path.join(output_dir, "master_cotizacion_simple.docx"), "SIMPLE")
create_informe_apa_master(os.path.join(output_dir, "master_informe_ejecutivo_apa.docx"))
create_informe_tecnico_master(os.path.join(output_dir, "master_informe_tecnico.docx"))
create_proyecto_pmi_master(os.path.join(output_dir, "master_proyecto_complejo_pmi.docx"), "COMPLEJO")
create_proyecto_pmi_master(os.path.join(output_dir, "master_proyecto_simple.docx"), "SIMPLE")

print("✅ 6 Master Documents created successfully.")
